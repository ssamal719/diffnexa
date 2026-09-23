"""Build the Excel Compare golden workbooks.

Run from apps/engine with the dev extras installed and LibreOffice available:

    python ../../golden/excel-pairs/_build/build_pairs.py

Each pair folder gets before.xlsx and after.xlsx. Workbooks are written with
openpyxl and then opened, calculated and saved by LibreOffice Calc, so every
formula carries its stored result exactly as a file saved by Excel does. The
expected.yaml beside them is written by hand and never generated here: it is
the human statement of what really changed.

The workbooks are realistic: a small company workbook (Summary, Pricing,
Employees, Settings) with currency, percentage and date formats, formulas that
depend on other sheets, and hyperlinks; and a 3,000-row inventory. Each pair
starts from one of them and makes only the edit its name describes. Pairs whose
expected.yaml states an expected error hold a damaged or unsupported file.

Formulas are written the way Excel keeps them after an edit: "{r}" is the
formula's own row and "{last}" the last pricing row, so inserting a row moves
and widens the formulas exactly as Excel would.
"""

from __future__ import annotations

import copy
import datetime as dt
import io
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
D = dt.date

MONEY = '"$"#,##0.00'
RUPEES = '"₹"#,##0'
EUROS = "[$€-407]#,##0.00"
PERCENT = "0%"
DATE = "d mmm yyyy"
LONG_DATE = "d mmmm yyyy"
ISO_DATE = "yyyy-mm-dd"

TOTAL = "=B{r}*C{r}*(1-D{r})"


def company() -> dict:
    return {
        "order": ["Summary", "Pricing", "Employees", "Settings"],
        "sheets": {
            "Summary": {
                "rows": [
                    ["Company Plan Summary"],
                    [],
                    ["Total monthly revenue", "=SUM(Pricing!E2:E{last})"],
                    ["Plans offered", "=COUNTA(Pricing!A2:A{last})"],
                    ["Average price", "=AVERAGE(Pricing!B2:B{last})"],
                    ["Report date", D(2026, 9, 30)],
                ],
                "formats": {"B3": MONEY, "B5": MONEY, "B6": LONG_DATE},
            },
            "Pricing": {
                "rows": [
                    ["Plan", "Price", "Seats", "Discount", "Total", "Renewal"],
                    ["Starter", 99, 1, 0, TOTAL, D(2026, 9, 30)],
                    ["Pro", 299, 3, 0.1, TOTAL, D(2026, 9, 30)],
                    ["Team", 499, 10, 0.15, TOTAL, D(2026, 10, 15)],
                    ["Enterprise", 1299, 25, 0.2, TOTAL, D(2026, 12, 31)],
                ],
                "columns": {"B": MONEY, "D": PERCENT, "E": MONEY, "F": DATE},
                "links": {"A2": "https://example.com/plans/starter"},
            },
            "Employees": {
                "rows": [
                    ["Name", "Role", "Start date", "Salary", "Location"],
                    ["John Carter", "Engineering Lead", D(2021, 4, 1), 2400000, "Bengaluru"],
                    ["Asha Rao", "Data Analyst", D(2023, 7, 17), 1150000, "Pune"],
                    ["Mei Lin", "Designer", D(2022, 1, 10), 1300000, None],
                    ["Omar Haddad", "Support Manager", D(2020, 11, 2), 1600000, "Hyderabad"],
                    ["Lena Novak", "Accountant", D(2024, 2, 19), 980000, "Remote"],
                ],
                "columns": {"C": ISO_DATE, "D": RUPEES},
            },
            "Settings": {
                "rows": [
                    ["Setting", "Value"],
                    ["Currency", "USD"],
                    ["Tax rate", 0.18],
                    ["Fiscal year start", D(2026, 4, 1)],
                    ["Support", "support@example.com"],
                ],
                "formats": {"B3": PERCENT, "B4": LONG_DATE},
                "links": {"B5": "mailto:support@example.com"},
            },
        },
    }


def inventory(rows: int = 3000) -> dict:
    data = [["SKU", "Product", "Category", "Unit price", "Stock", "Reorder at", "Warehouse", "Last counted"]]
    categories = ["Tools", "Garden", "Kitchen", "Office", "Outdoor", "Lighting"]
    for i in range(1, rows + 1):
        data.append(
            [
                f"SKU-{i:05d}",
                f"Item {i:04d}",
                categories[i % len(categories)],
                round(5 + (i * 37 % 900) / 10, 2),
                (i * 13) % 250,
                20 + i % 15,
                f"WH-{1 + i % 4}",
                D(2026, 1, 1) + dt.timedelta(days=i % 240),
            ]
        )
    return {
        "order": ["Products"],
        "sheets": {"Products": {"rows": data, "columns": {"D": EUROS, "H": DATE}}},
    }


def render(spec: dict) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    last = len(spec["sheets"].get("Pricing", {}).get("rows", [])) or 1
    for name in spec["order"]:
        sheet_spec = spec["sheets"][name]
        sheet = book.create_sheet(name)
        for r, row in enumerate(sheet_spec["rows"], start=1):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                if isinstance(value, str) and value.startswith("="):
                    value = value.format(r=r, last=last)
                sheet.cell(row=r, column=c, value=value)
        for column, number_format in sheet_spec.get("columns", {}).items():
            for r in range(2, len(sheet_spec["rows"]) + 1):
                sheet[f"{column}{r}"].number_format = number_format
        for ref, number_format in sheet_spec.get("formats", {}).items():
            sheet[ref].number_format = number_format
        for ref, target in sheet_spec.get("links", {}).items():
            sheet[ref].hyperlink = target
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def saved_by_calc(data: bytes) -> bytes:
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "book.xlsx"
        source.write_bytes(data)
        out = Path(directory) / "out"
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{directory}/profile",
                "--headless",
                "--calc",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out),
                str(source),
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )
        return (out / "book.xlsx").read_bytes()


# ---------------------------------------------------------------- edits


def rows(spec: dict, sheet: str) -> list[list]:
    return spec["sheets"][sheet]["rows"]


def cell(spec: dict, sheet: str, ref: str, value) -> None:
    col = openpyxl.utils.column_index_from_string("".join(ch for ch in ref if ch.isalpha()))
    row = int("".join(ch for ch in ref if ch.isdigit()))
    grid = rows(spec, sheet)
    while len(grid) < row:
        grid.append([])
    line = grid[row - 1]
    while len(line) < col:
        line.append(None)
    line[col - 1] = value


PAIRS: dict[str, tuple] = {}


def pair(name: str, base, edit=None) -> None:
    before = base()
    after = copy.deepcopy(before)
    if edit:
        edit(after)
    PAIRS[name] = ("workbooks", before, after)


def broken(name: str, before: bytes, after: bytes) -> None:
    PAIRS[name] = ("files", before, after)


pair("no-change", company)
pair("text-change", company, lambda s: cell(s, "Employees", "A2", "Jonathan Carter"))
pair("number-change", company, lambda s: cell(s, "Pricing", "B3", 349))
pair("date-change", company, lambda s: cell(s, "Pricing", "F4", D(2026, 11, 1)))
pair("formula-change", company, lambda s: cell(s, "Pricing", "E5", "=B{r}*C{r}*(1-D{r})*0.95"))
pair("added-cell", company, lambda s: cell(s, "Employees", "E4", "Chennai"))
pair("removed-cell", company, lambda s: cell(s, "Employees", "E6", None))
pair(
    "added-row",
    company,
    lambda s: rows(s, "Pricing").insert(3, ["Growth", 199, 5, 0.05, TOTAL, D(2026, 9, 30)]),
)
pair("removed-row", company, lambda s: rows(s, "Employees").pop(3))
pair(
    "added-column",
    company,
    lambda s: (
        [
            row.insert(2, value)
            for row, value in zip(
                rows(s, "Employees"),
                ["Department", "Engineering", "Analytics", "Design", "Support", "Finance"],
                strict=True,
            )
        ]
        and s["sheets"]["Employees"].update(columns={"D": ISO_DATE, "E": RUPEES})
    ),
)
pair(
    "removed-column",
    company,
    lambda s: [row.pop(4) for row in rows(s, "Employees") if len(row) > 4],
)


def _add_sheet(s):
    s["order"].append("Regions")
    s["sheets"]["Regions"] = {"rows": [["Region", "Target"], ["North", 120000], ["South", 95000]]}


pair("added-sheet", company, _add_sheet)
pair("removed-sheet", company, lambda s: (s["order"].remove("Settings"), s["sheets"].pop("Settings")))


def _rename(s):
    s["order"][s["order"].index("Employees")] = "Staff"
    s["sheets"]["Staff"] = s["sheets"].pop("Employees")


pair("sheet-rename", company, _rename)


def _move_sheet(s):
    s["order"] = ["Summary", "Settings", "Pricing", "Employees"]


pair("sheet-order", company, _move_sheet)
pair(
    "hyperlink-change",
    company,
    lambda s: s["sheets"]["Pricing"]["links"].update(A2="https://example.com/plans/basic"),
)


def _moved_row(s):
    grid = rows(s, "Employees")
    grid.append(grid.pop(1))


pair("moved-row", company, _moved_row)


def _blank_rows(s):
    grid = rows(s, "Employees")
    grid[3:3] = [[], []]


pair("blank-rows-inserted", company, _blank_rows)
pair("formula-result", company, lambda s: cell(s, "Pricing", "C4", 12))


def _dates(s):
    cell(s, "Summary", "B6", D(2026, 10, 31))
    cell(s, "Employees", "C3", D(2023, 8, 1))
    cell(s, "Settings", "B4", D(2026, 7, 1))


pair("dates", company, _dates)


def _currency(s):
    cell(s, "Employees", "D5", 1750000)
    cell(s, "Pricing", "B2", 109)


pair("currency", company, _currency)
pair("percentages", company, lambda s: (cell(s, "Settings", "B3", 0.2), cell(s, "Pricing", "D4", 0.2)))
def company_with_review() -> dict:
    spec = company()
    spec["sheets"]["Settings"]["rows"].append(["Next review", "30 September 2026"])  # typed as text
    return spec


pair("text-dates", company_with_review, lambda s: cell(s, "Settings", "B6", "15 October 2026"))


def _format_only(s):
    s["sheets"]["Pricing"]["columns"]["B"] = "#,##0"
    s["sheets"]["Pricing"]["columns"]["F"] = ISO_DATE


pair("format-only", company, _format_only)


def _multiple_sheets(s):
    cell(s, "Summary", "A1", "Company Plan Summary — Q4")
    cell(s, "Pricing", "C3", 4)
    cell(s, "Employees", "B3", "Senior Data Analyst")
    cell(s, "Settings", "B2", "EUR")


pair("multiple-sheets", company, _multiple_sheets)


def _mixed(s):
    cell(s, "Pricing", "B3", 349)
    rows(s, "Pricing").insert(3, ["Growth", 199, 5, 0.05, TOTAL, D(2026, 9, 30)])
    cell(s, "Pricing", "F5", D(2026, 11, 1))
    cell(s, "Employees", "A2", "Jonathan Carter")
    rows(s, "Employees").pop(4)
    s["sheets"]["Pricing"]["links"]["A2"] = "https://example.com/plans/basic"
    _add_sheet(s)


pair("multiple-changes", company, _mixed)


def _inventory_changes(s):
    grid = rows(s, "Products")
    grid[1500][3] = 99.99  # SKU-01500 unit price
    grid[2200][4] = 0  # SKU-02200 stock
    grid.insert(2801, ["SKU-02800A", "Item 2800A", "Tools", 12.5, 40, 25, "WH-2", D(2026, 9, 30)])


pair("large-workbook", inventory, _inventory_changes)


def _packaged(
    sheet_xml: str,
    content_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    members=(),
) -> bytes:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/xl/workbook.xml" ContentType="{content_type}"/></Types>'
        ),
        "_rels/.rels": (
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<?xml version="1.0"?><workbook xmlns="{main_ns}" xmlns:r="{rel}"><sheets>'
            '<sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": sheet_xml,
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
        for name, data in members:
            archive.writestr(name, data)
    return buffer.getvalue()


def build_broken(good: bytes) -> None:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    broken("malformed-workbook", good, good[: len(good) // 3])
    broken("unsupported-xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 4096, good)
    broken(
        "unsupported-xlsm",
        good,
        _packaged(
            f'<worksheet xmlns="{main_ns}"><sheetData/></worksheet>',
            "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            [("xl/vbaProject.bin", b"\x00macro")],
        ),
    )
    padding = " " * (70 * 1024 * 1024)
    broken(
        "security-zip-bomb",
        good,
        _packaged(
            f'<worksheet xmlns="{main_ns}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is>'
            f'<t xml:space="preserve">{padding}</t></is></c></row></sheetData></worksheet>'
        ),
    )
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        traversal = _packaged(
            f'<worksheet xmlns="{main_ns}"><sheetData/></worksheet>', members=[("../../outside.txt", b"x")]
        )
    broken("security-path-traversal", traversal, good)
    broken(
        "security-external-entity",
        good,
        _packaged(
            '<?xml version="1.0"?><!DOCTYPE w [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
            f'<worksheet xmlns="{main_ns}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&x;</t>'
            "</is></c></row></sheetData></worksheet>"
        ),
    )


def main() -> int:
    if shutil.which("soffice") is None:
        print("LibreOffice (soffice) is required to store formula results.")
        return 1
    only = set(sys.argv[1:])
    cache: dict[int, bytes] = {}

    def made(spec: dict) -> bytes:
        key = hash(repr(spec))
        if key not in cache:
            cache[key] = saved_by_calc(render(spec))
        return cache[key]

    build_broken(made(company()))
    for name, (kind, before, after) in PAIRS.items():
        if only and name not in only:
            continue
        folder = ROOT / name
        folder.mkdir(parents=True, exist_ok=True)
        if kind == "workbooks":
            (folder / "before.xlsx").write_bytes(made(before))
            (folder / "after.xlsx").write_bytes(made(after))
        else:
            (folder / "before.xlsx").write_bytes(before)
            (folder / "after.xlsx").write_bytes(after)
        if not (folder / "expected.yaml").exists():
            print(f"{name}: write expected.yaml by hand")
    print(f"Built {len(PAIRS)} pairs in {ROOT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
