"""Building real .docx files for tests.

Two builders:

* `word(...)` makes an ordinary document with python-docx, so headings, lists
  and tables carry Word's real styles and numbering definitions, exactly as a
  document saved from Word does.
* `package(...)` assembles a .docx by hand from XML strings, for the cases
  python-docx will not write: tracked changes, hostile archives, macro-enabled
  content types, entity declarations.

Nothing here converts a document to plain text: every test goes through the
real reader.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable

import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DOCX_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def word(build: Callable[[docx.document.Document], None]) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def add_hyperlink(paragraph, url: str, text: str) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def table(document, rows: list[tuple[str, ...]]) -> None:
    grid = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            grid.cell(i, j).text = value


def body(xml: str) -> str:
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>{xml}</w:body></w:document>'
    )


def package(
    document_xml: str | bytes,
    *,
    content_type: str = DOCX_MAIN,
    document_rels: str = "",
    extra: dict[str, bytes | str] | None = None,
    members: list[tuple[str, bytes | str]] | None = None,
    statistics: tuple[int, int] | None = None,
) -> bytes:
    """A minimal but genuine .docx: content types, package rels, the document.

    `statistics` adds docProps/app.xml with (pages, words), as Word writes it.
    """
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/word/document.xml" ContentType="{content_type}"/>'
        "</Types>"
    )
    package_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/officeDocument" Target="word/document.xml"/>'
        + (
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/extended-properties" Target="docProps/app.xml"/>'
            if statistics
            else ""
        )
        + "</Relationships>"
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f"{document_rels}</Relationships>"
    )
    entries: list[tuple[str, bytes | str]] = [
        ("[Content_Types].xml", content_types),
        ("_rels/.rels", package_rels),
        ("word/document.xml", document_xml),
        ("word/_rels/document.xml.rels", doc_rels),
    ]
    if statistics:
        pages, words = statistics
        entries.append(
            (
                "docProps/app.xml",
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                f"<Application>Microsoft Office Word</Application><Pages>{pages}</Pages>"
                f"<Words>{words}</Words></Properties>",
            )
        )
    entries += list((extra or {}).items())
    entries += members or []
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return buffer.getvalue()


def paragraph(text: str) -> str:
    return f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'
