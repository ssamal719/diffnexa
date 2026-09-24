"""Builds the 100+ page Word documents used to test DOCX page locations.

Run from apps/engine with LibreOffice installed:

    python tests/fixtures/docx_pages/build.py

What it does, and why the result is a fair test:

1. Writes a long document with python-docx: headings, paragraphs, bulleted
   items, tables and two hard page breaks. Every word is unique ("p0123w045"),
   so a word's page can be found in a rendered PDF without guessing.
2. Renders it with LibreOffice — an independent layout engine — and reads
   which page every word landed on.
3. Records that layout in the file the way Microsoft Word does on save: a
   <w:lastRenderedPageBreak/> at the exact word where each page began, and the
   page and word counts in docProps/app.xml. (LibreOffice renders but does not
   write these markers; Word does.)
4. Makes a revised document with changes placed on chosen pages of the
   original rendering — a word, a number, an added and a removed paragraph, a
   change in the second half of a paragraph that runs across a page break, two
   changes on one page, and a table cell — and records its layout the same way.
5. Writes expected.json: for every change, the page the rendered PDF shows it
   on, in each document. That is the test's oracle. It comes from the PDFs,
   never from counting paragraphs.

The committed .docx files are what the tests read; LibreOffice is only needed
to rebuild them.
"""

from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
import pdfplumber
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = Path(__file__).parent
TOKEN = re.compile(r"^[a-z]+\d{4}[a-z]\d{3}$")


# ---------------------------------------------------------------- content


def word(prefix: str, block: int, index: int) -> str:
    return f"{prefix}{block:04d}w{index:03d}"


def base_blocks() -> list[dict]:
    """The original document, as a list of blocks with their words."""
    blocks: list[dict] = []
    for i in range(1, 331):
        if i % 25 == 1:
            blocks.append({"kind": "heading", "id": f"h{i}", "words": [word("h", i, j) for j in range(1, 4)]})
        if i in (120, 240):
            blocks.append({"kind": "pagebreak", "id": f"pb{i}"})
        count = 90 + (i * 37) % 80  # 90 to 169 words: many paragraphs run across a page break
        words = [word("p", i, j) for j in range(1, count + 1)]
        if i % 3 == 0:
            # A quantity: a number with unique words on both sides of it.
            words[20:20] = ["Quantity", str(100 + i), "units"]
        blocks.append({"kind": "para", "id": f"p{i}", "words": words})
        if i % 17 == 0:
            blocks.append({"kind": "bullet", "id": f"b{i}", "words": [word("b", i, j) for j in range(1, 12)]})
        if i % 40 == 0:
            rows = [[[word("t", i, r * 10 + c)] for c in range(1, 4)] for r in range(1, 4)]
            blocks.append({"kind": "table", "id": f"t{i}", "rows": rows})
    return blocks


def write_docx(blocks: list[dict], path: Path) -> None:
    document = docx.Document()
    for block in blocks:
        if block["kind"] == "heading":
            document.add_heading(" ".join(block["words"]), level=1)
        elif block["kind"] == "para":
            document.add_paragraph(" ".join(block["words"]))
        elif block["kind"] == "bullet":
            document.add_paragraph(" ".join(block["words"]), style="List Bullet")
        elif block["kind"] == "pagebreak":
            document.add_page_break()
        elif block["kind"] == "table":
            rows = block["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell_words in enumerate(row):
                    table.cell(r, c).text = " ".join(cell_words)
    document.save(path)


# ---------------------------------------------------------------- rendering


def render(path: Path, work: Path) -> Path:
    out = work / "pdf"
    out.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
        check=True,
        capture_output=True,
        timeout=600,
    )
    return out / (path.stem + ".pdf")


def token_pages(pdf: Path) -> tuple[dict[str, int], int]:
    pages: dict[str, int] = {}
    with pdfplumber.open(pdf) as document:
        for number, page in enumerate(document.pages, start=1):
            for item in page.extract_words():
                text = item["text"]
                if TOKEN.match(text):
                    if text in pages:
                        raise SystemExit(f"{text} appears twice in the rendering")
                    pages[text] = number
        return pages, len(document.pages)


def page_of_word(words: list[str], index: int, pages: dict[str, int]) -> int:
    """A word's page. Numbers and ordinary words take the page of the unique
    word before them, which must agree with the unique word after them."""
    if TOKEN.match(words[index]):
        return pages[words[index]]
    before = next(pages[w] for w in reversed(words[:index]) if TOKEN.match(w))
    after = next((pages[w] for w in words[index + 1 :] if TOKEN.match(w)), before)
    if before != after:
        raise SystemExit(f"cannot place {words[index]!r}: a page break falls beside it")
    return before


# ---------------------------------------------------------------- recording the layout like Word


def _marker() -> OxmlElement:
    return OxmlElement("w:lastRenderedPageBreak")


def _split_paragraph(paragraph, words: list[str], starts: list[int]) -> None:
    """Rewrite a one-run paragraph as runs that each begin where a page begins."""
    runs = paragraph.runs
    style = copy.deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    for run in runs:
        run._r.getparent().remove(run._r)
    cuts = sorted({0, *starts, len(words)})
    for a, b in zip(cuts, cuts[1:], strict=False):
        text = " ".join(words[a:b]) + (" " if b < len(words) else "")
        run = paragraph.add_run(text)
        if style is not None:
            run._r.insert(0, copy.deepcopy(style))
        if a in starts:
            # Word puts the marker inside the run, before the text that begins the page.
            run._r.insert(1 if run._r.rPr is not None else 0, _marker())


def record_layout(blocks: list[dict], source: Path, target: Path, pages: dict[str, int], total: int) -> None:
    """Write Word's rendered-page markers and statistics into a copy of `source`."""
    document = docx.Document(source)
    items = [i for i in document.element.body.iterchildren() if i.tag in (qn("w:p"), qn("w:tbl"))]

    previous_page = 1
    after_hard_break = False
    first_hard_break_seen = False
    word_total = 0
    for block, element in zip(blocks, items, strict=True):
        if block["kind"] == "pagebreak":
            after_hard_break = True
            continue
        if block["kind"] == "table":
            table = Table(element, document._body)
            for r, row in enumerate(block["rows"]):
                row_page = page_of_word(row[0], 0, pages)
                for c, cell_words in enumerate(row):
                    word_total += len(cell_words)
                    if row_page > previous_page:
                        # Word marks a row that starts a new page in every cell.
                        _split_paragraph(table.cell(r, c).paragraphs[0], cell_words, [0])
                previous_page = max(previous_page, row_page)
            continue
        words = block["words"]
        word_total += len(words)
        starts: list[int] = []
        for index in range(len(words)):
            page = page_of_word(words, index, pages)
            if page > previous_page:
                if page != previous_page + 1:
                    raise SystemExit(f"{block['id']}: skipped a page")
                # After a hard page break Word may or may not also write its own
                # marker; the first break here gets one and the second does not,
                # so both are exercised.
                if not (after_hard_break and first_hard_break_seen):
                    starts.append(index)
                if after_hard_break:
                    first_hard_break_seen = True
                previous_page = page
                after_hard_break = False
        if starts:
            _split_paragraph(Paragraph(element, document._body), words, starts)
    if previous_page != total:
        raise SystemExit(f"layout covers {previous_page} pages but the rendering has {total}")
    document.save(target)
    _write_statistics(target, total, word_total)


def _write_statistics(path: Path, pages: int, words: int) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/app.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<Pages>\d+</Pages>", f"<Pages>{pages}</Pages>", text)
                text = re.sub(r"<Words>\d+</Words>", f"<Words>{words}</Words>", text)
                text = text.replace(
                    "Microsoft Macintosh Word", "DiffNexa test builder (layout from LibreOffice)"
                )
                data = text.encode("utf-8")
            out.writestr(info, data)
    temporary.replace(path)


# ---------------------------------------------------------------- the changes


def fully_on(block: dict, page: int, pages: dict[str, int]) -> bool:
    words = block.get("words") or []
    return bool(words) and all(page_of_word(words, i, pages) == page for i in range(len(words)))


def pick(
    blocks: list[dict], pages: dict[str, int], page: int, kind: str = "para", quantity: bool = False
) -> dict:
    for block in blocks:
        if (
            block["kind"] == kind
            and fully_on(block, page, pages)
            and ("Quantity" in block["words"]) == quantity
        ):
            return block
    raise SystemExit(f"no {kind} fully on page {page}")


def spanning(blocks: list[dict], pages: dict[str, int], near: int) -> tuple[dict, int]:
    for block in blocks:
        if block["kind"] != "para":
            continue
        words = block["words"]
        first, last = page_of_word(words, 0, pages), page_of_word(words, len(words) - 1, pages)
        if first >= near and last == first + 1:
            split = next(i for i in range(len(words)) if page_of_word(words, i, pages) == last)
            if split + 3 < len(words):
                return block, split
    raise SystemExit("no paragraph spans a page break")


def plan_changes(blocks: list[dict], pages: dict[str, int]) -> tuple[list[dict], list[dict]]:
    revised = copy.deepcopy(blocks)
    by_id = {block["id"]: block for block in revised}
    changes: list[dict] = []

    def replace_word(block_id: str, index: int, new: str, change_id: str, what: str) -> None:
        block = by_id[block_id]
        old = block["words"][index]
        block["words"][index] = new
        changes.append({"id": change_id, "what": what, "old": old, "new": new})

    a = pick(blocks, pages, 10)
    replace_word(a["id"], 5, "amended0010x001", "text-page-10", "a word replaced")
    b = pick(blocks, pages, 25, quantity=True)
    q = b["words"].index("Quantity") + 1
    replace_word(b["id"], q, str(int(b["words"][q]) + 27), "number-page-25", "a quantity changed")
    c = pick(blocks, pages, 50)
    added = {"kind": "para", "id": "added50", "words": [word("added", 50, j) for j in range(1, 41)]}
    revised.insert(revised.index(by_id[c["id"]]) + 1, added)
    changes.append(
        {"id": "added-page-50", "what": "a paragraph added", "old": None, "new": added["words"][0]}
    )
    d = pick(blocks, pages, 75)
    revised.remove(by_id[d["id"]])
    changes.append(
        {"id": "removed-page-75", "what": "a paragraph removed", "old": d["words"][0], "new": None}
    )
    e, split = spanning(blocks, pages, 60)
    replace_word(
        e["id"],
        split + 2,
        "amended0060x002",
        "second-half-of-a-split-paragraph",
        "a word after a page break inside a paragraph",
    )
    # Two changes on one page: the first page from 100 on that holds two whole paragraphs.
    same_page = next(
        page
        for page in range(100, 200)
        if sum(1 for block in blocks if block["kind"] == "para" and fully_on(block, page, pages)) >= 2
    )
    f, g = [block for block in blocks if block["kind"] == "para" and fully_on(block, same_page, pages)][:2]
    replace_word(f["id"], 3, "amended0100x003", "same-page-first", "one of two changes on one page")
    replace_word(g["id"], 7, "amended0100x005", "same-page-second", "the other change on that page")
    table = next(
        block
        for block in blocks
        if block["kind"] == "table" and page_of_word(block["rows"][1][1], 0, pages) >= 70
    )
    old_cell = table["rows"][1][1][0]
    by_id[table["id"]]["rows"][1][1] = ["amended0070x004"]
    changes.append(
        {"id": "table-cell", "what": "a table cell changed", "old": old_cell, "new": "amended0070x004"}
    )
    return revised, changes


# ---------------------------------------------------------------- build


def main() -> None:
    if shutil.which("soffice") is None:
        sys.exit("LibreOffice (soffice) is needed to rebuild these fixtures.")
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        original_blocks = base_blocks()
        write_docx(original_blocks, work / "original-plain.docx")
        original_pages, original_total = token_pages(render(work / "original-plain.docx", work))

        revised_blocks, changes = plan_changes(original_blocks, original_pages)
        write_docx(revised_blocks, work / "revised-plain.docx")
        revised_pages, revised_total = token_pages(render(work / "revised-plain.docx", work))

        record_layout(
            original_blocks,
            work / "original-plain.docx",
            HERE / "original.docx",
            original_pages,
            original_total,
        )
        record_layout(
            revised_blocks, work / "revised-plain.docx", HERE / "revised.docx", revised_pages, revised_total
        )

        # The markers must not change LibreOffice's own layout of the files.
        for name, expected in (("original", original_pages), ("revised", revised_pages)):
            again, _ = token_pages(render(HERE / f"{name}.docx", work))
            if again != expected:
                sys.exit(f"{name}.docx renders differently once its layout is recorded")

        for change in changes:
            change["oldPage"] = (
                original_pages.get(change["old"]) if change["old"] and TOKEN.match(change["old"]) else None
            )
            change["newPage"] = (
                revised_pages.get(change["new"]) if change["new"] and TOKEN.match(change["new"]) else None
            )
            if change["old"] and not TOKEN.match(change["old"]):  # a number: placed by the words around it
                block = next(b for b in original_blocks if change["old"] in (b.get("words") or []))
                change["oldPage"] = page_of_word(
                    block["words"], block["words"].index(change["old"]), original_pages
                )
            if change["new"] and not TOKEN.match(change["new"]):
                block = next(b for b in revised_blocks if change["new"] in (b.get("words") or []))
                change["newPage"] = page_of_word(
                    block["words"], block["words"].index(change["new"]), revised_pages
                )

        oracle = {
            "note": "Pages read from LibreOffice renderings of the plain documents; see build.py.",
            "pages": {"original": original_total, "revised": revised_total},
            "changes": changes,
        }
        (HERE / "expected.json").write_text(json.dumps(oracle, indent=2) + "\n")
        print(json.dumps(oracle["pages"]), len(changes), "changes")
        for change in changes:
            print(f"  {change['id']:36} original p.{change['oldPage']}  revised p.{change['newPage']}")


if __name__ == "__main__":
    main()
