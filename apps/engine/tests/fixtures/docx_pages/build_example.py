"""Builds the example Word documents behind DOCX Compare's "Try Example".

Run from apps/engine with LibreOffice installed:

    python tests/fixtures/docx_pages/build_example.py

The two documents are a short, fictional services agreement and a revision of
it. They are ordinary documents run through the real comparison — the example
shows genuine results, never prepared ones. Like the long test documents (see
build.py), each is laid out by LibreOffice and that layout is recorded in the
file the way Microsoft Word records its own, so the example shows page numbers.

The revision changes a fee in a table, a notice period, a start date and a
heading; adds one clause and removes another; and makes two edits the Ignore
options act on: a capitalisation-only edit (ignored by default) and a
punctuation-only edit (reported unless punctuation is ignored).
"""

from __future__ import annotations

import sys
import tempfile
from difflib import SequenceMatcher
from pathlib import Path

import docx
import pdfplumber
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).parent))
from build import _split_paragraph, _write_statistics, render  # noqa: E402

OUT = Path(__file__).resolve().parents[5] / "apps" / "web" / "public" / "examples"

PARTIES = "Harbourline Design Studio (the Studio) and Example Client Ltd (the Client)"


def agreement(revised: bool) -> list[dict]:
    b: list[dict] = []

    def heading(text: str) -> None:
        b.append({"kind": "heading", "text": text})

    def para(text: str) -> None:
        b.append({"kind": "para", "text": text})

    def bullet(text: str) -> None:
        b.append({"kind": "bullet", "text": text})

    heading("Website Design Services Agreement")
    para(
        f"This agreement is made between {PARTIES}. It sets out the design services the Studio will provide, "
        "the fees the Client will pay, and how either party may end "
        "the arrangement. It is an example document "
        "written for DiffNexa; the parties and terms are fictional."
    )
    heading("1. Services")
    para(
        "The Studio will design and build a marketing website for the Client"
        + (" including" if revised else ", including")
        + " page layouts, a small component library, copy editing for up to twelve pages, and two rounds of "
        "revisions on each deliverable. The Studio will share work "
        "in progress every two weeks and will record "
        "decisions in a shared project log that both parties can read."
    )
    for item in (
        "Discovery workshop and written project brief.",
        "Page designs for desktop and mobile screens.",
        "Development of the approved designs and launch support.",
    ):
        bullet(item)
    para(
        "Work begins on "
        + ("1 April 2026" if revised else "1 March 2026")
        + " and is planned to take sixteen "
        "weeks. The timetable depends on the Client providing "
        "content and feedback within five working days of "
        "each request; delays on either side move the dates that follow by the same number of days."
    )
    for index in range(1, 5):
        para(
            f"Milestone {index}: the Studio will present the deliverables "
            "for this stage in a review meeting, "
            "explain the decisions behind them, and collect the Client's "
            "comments in the project log. Comments "
            "received after the review window closes will be scheduled "
            "into the next stage rather than reopening "
            "the current one, so that the overall timetable can be kept."
        )
    heading("2. Fees and payment")
    para(
        "The Client will pay a monthly retainer for the duration of "
        "the project, invoiced at the start of each "
        "month and payable within thirty days. Work outside the "
        "agreed services is charged at the day rate below "
        "once the Client has approved a written estimate."
    )
    b.append(
        {
            "kind": "table",
            "rows": [
                ["Item", "Amount"],
                ["Monthly retainer", "£2,750" if revised else "£2,500"],
                ["Additional day rate", "£650"],
                ["Hosting set-up (one-off)", "£300"],
            ],
        }
    )
    if not revised:
        para("Travel within the city is included; other travel will be reimbursed at cost with receipts.")
    heading("3. Intellectual property")
    para(
        "Once all fees are paid, the Client owns the final designs and the website code written for this "
        "project. The Studio keeps ownership of its general tools, "
        "templates and know-how, and grants the Client "
        "a permanent licence to use any of them that are built into the website."
    )
    heading("4. Data protection")
    para(
        "Each party will handle personal data it receives under this agreement lawfully and only for the "
        "purposes of the project. The Studio will keep project files in access-controlled storage and delete "
        "them within ninety days of the project ending, unless the Client asks for them to be returned first."
    )
    if revised:
        para(
            "The Studio will not appoint a subcontractor to process the Client's personal data without the "
            "Client's prior written approval, and will remain responsible for any subcontractor it appoints."
        )
    heading("5. Ending the agreement")
    para(
        "Either party may end this agreement by giving "
        + ("60" if revised else "30")
        + " days' written notice. The Client will pay for work completed up to the date the agreement ends, "
        "and the Studio will hand over all finished work and working files that the Client has paid for."
    )
    heading("6. Limitation of liability" if revised else "6. Liability")
    para(
        "Neither party is liable for indirect or consequential "
        "losses. Each party's total liability under this "
        "agreement is limited to the fees paid in the twelve months "
        "before the claim arose, except where the law "
        "does not allow liability to be limited."
    )
    heading("7. General")
    para(
        "This agreement is the whole agreement between the parties about the project. Changes to it must be "
        "agreed in writing by both parties. If any part of it cannot "
        "be enforced, the rest remains in effect. "
        + "Notices must be sent to the "
        + ("client's" if revised else "Client's")
        + " registered email address."
    )
    return b


def write(blocks: list[dict], path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Website Design Services Agreement (example)"
    for block in blocks:
        if block["kind"] == "heading":
            level = 0 if block["text"].startswith("Website") else 1
            document.add_heading(block["text"], level=level)
        elif block["kind"] == "para":
            document.add_paragraph(block["text"])
        elif block["kind"] == "bullet":
            document.add_paragraph(block["text"], style="List Bullet")
        elif block["kind"] == "table":
            rows = block["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
    document.save(path)


def word_pages(blocks: list[dict], pdf: Path) -> tuple[list[list[int]], int]:
    """The page of every word, found by aligning the document's words with the rendering's."""
    rendered: list[tuple[str, int]] = []
    with pdfplumber.open(pdf) as document:
        total = len(document.pages)
        for number, page in enumerate(document.pages, start=1):
            rendered += [(item["text"], number) for item in page.extract_words()]
    sequence: list[tuple[int, int, str]] = []  # (block, word, text)
    for index, block in enumerate(blocks):
        if block["kind"] == "table":
            for row in block["rows"]:
                for text in row:
                    for position, token in enumerate(text.split()):
                        sequence.append((index, position, token))
        else:
            for position, token in enumerate(block["text"].split()):
                sequence.append((index, position, token))
    matcher = SequenceMatcher(a=[t for _, _, t in sequence], b=[t for t, _ in rendered], autojunk=False)
    pages = [0] * len(sequence)
    for a, b_, size in matcher.get_matching_blocks():
        for offset in range(size):
            pages[a + offset] = rendered[b_ + offset][1]
    last = 1
    for i, page in enumerate(pages):  # an unmatched word stays on the page before it
        pages[i] = page or last
        last = pages[i]
    per_block: list[list[int]] = [[] for _ in blocks]
    for (block, _, _), page in zip(sequence, pages, strict=True):
        per_block[block].append(page)
    return per_block, total


def record(blocks: list[dict], source: Path, target: Path, pages: list[list[int]], total: int) -> None:
    document = docx.Document(source)
    items = [i for i in document.element.body.iterchildren() if i.tag in (qn("w:p"), qn("w:tbl"))]
    previous = 1
    words = 0
    for block, element, block_pages in zip(blocks, items, pages, strict=True):
        if block["kind"] == "table":
            table = Table(element, document._body)
            cursor = 0
            for r, row in enumerate(block["rows"]):
                row_page = block_pages[cursor]
                for c, text in enumerate(row):
                    count = len(text.split())
                    words += count
                    if row_page > previous:
                        _split_paragraph(table.cell(r, c).paragraphs[0], text.split(), [0])
                    cursor += count
                previous = max(previous, row_page)
            continue
        tokens = block["text"].split()
        words += len(tokens)
        starts = []
        for index, page in enumerate(block_pages):
            if page > previous:
                starts.append(index)
                previous = page
        if starts:
            _split_paragraph(Paragraph(element, document._body), tokens, starts)
    if previous != total:
        raise SystemExit(f"layout covers {previous} pages but the rendering has {total}")
    document.save(target)
    _write_statistics(target, total, words)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        for name, revised in (("original", False), ("revised", True)):
            blocks = agreement(revised)
            plain = work / f"{name}.docx"
            write(blocks, plain)
            pages, total = word_pages(blocks, render(plain, work))
            target = OUT / f"service-agreement-{name}.docx"
            record(blocks, plain, target, pages, total)
            print(target.name, total, "pages")


if __name__ == "__main__":
    main()
